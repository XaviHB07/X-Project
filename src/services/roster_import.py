"""Lectura y validación del Excel de la lista de estudiantes (HU-C1).

Este módulo es PURO: recibe los bytes de un `.xlsx` y devuelve filas ya
validadas más una lista de errores por fila. No toca la base de datos ni
FastAPI, así que se prueba sin infraestructura (ver
`tests/test_roster_import.py`).

Formato esperado (la primera fila no vacía es el encabezado):

    | codigo   | nombre        |
    | 2024-001 | Ana Torres    |
    | 2024-002 | Luis Pérez    |

Encabezados aceptados (sin distinguir mayúsculas, tildes, espacios ni
guiones bajos):

    - código:  codigo, código, matricula, matrícula, external_ref, dni, id
    - nombre:  nombre, nombres, nombre completo, apellidos y nombres,
               estudiante, display_name

Dos niveles de problema, a propósito:

    * Errores de ESTRUCTURA (archivo que no es .xlsx, vacío, sin las
      columnas requeridas, demasiado grande): `RosterFormatError`. No se
      puede interpretar nada, así que no se importa nada.
    * Errores de REGISTRO (fila sin nombre, código repetido, texto
      demasiado largo): se devuelven en `RosterParseResult.errors` con
      su número de fila de Excel; las demás filas siguen siendo válidas.
"""

from __future__ import annotations

import io
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Set

from openpyxl import load_workbook

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_FIELD_LENGTH = 255

_CODE_HEADERS = {"codigo", "matricula", "externalref", "dni", "id", "codigodematricula"}
_NAME_HEADERS = {
    "nombre", "nombres", "nombrecompleto", "apellidosynombres", "estudiante", "displayname",
}


class RosterFormatError(ValueError):
    """El archivo no se puede interpretar como lista de estudiantes."""


@dataclass(frozen=True)
class RosterRow:
    row_number: int  # número de fila tal como lo ve el docente en Excel
    external_ref: str
    display_name: str


@dataclass(frozen=True)
class RowError:
    row_number: int
    message: str


@dataclass
class RosterParseResult:
    rows: List[RosterRow] = field(default_factory=list)
    errors: List[RowError] = field(default_factory=list)
    total_rows: int = 0  # filas de datos no vacías leídas


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return "".join(ch for ch in text if ch.isalnum())


def _cell_text(value: Any) -> str:
    """Convierte una celda a texto limpio. Los códigos numéricos de Excel
    (p. ej. 20240001) llegan como int/float: se pasan a texto sin el
    ".0" final para no cambiar el código del estudiante."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def parse_roster_xlsx(data: bytes) -> RosterParseResult:
    if not data:
        raise RosterFormatError("El archivo está vacío.")
    if len(data) > MAX_FILE_BYTES:
        raise RosterFormatError(
            f"El archivo supera el máximo permitido ({MAX_FILE_BYTES // (1024 * 1024)} MB)."
        )

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza tipos distintos según el daño del archivo
        raise RosterFormatError(
            "No se pudo leer el archivo. Debe ser un Excel en formato .xlsx."
        ) from exc

    try:
        sheet = workbook.active
        if sheet is None:
            raise RosterFormatError("El Excel no tiene hojas.")
        row_iter = sheet.iter_rows(values_only=True)

        header_cells: Optional[tuple] = None
        header_row_number = 0
        for header_row_number, cells in enumerate(row_iter, start=1):
            if any(_cell_text(c) for c in cells):
                header_cells = cells
                break
        if header_cells is None:
            raise RosterFormatError("El Excel no tiene datos.")

        code_col = name_col = None
        for index, cell in enumerate(header_cells):
            key = _normalize_header(cell)
            if code_col is None and key in _CODE_HEADERS:
                code_col = index
            elif name_col is None and key in _NAME_HEADERS:
                name_col = index
        missing = []
        if code_col is None:
            missing.append("código (codigo / matrícula)")
        if name_col is None:
            missing.append("nombre (nombre / apellidos y nombres)")
        if missing:
            raise RosterFormatError(
                "Faltan columnas obligatorias en el encabezado: " + " y ".join(missing) + "."
            )

        result = RosterParseResult()
        seen: Dict[str, int] = {}
        current_row = header_row_number
        for current_row, cells in enumerate(row_iter, start=header_row_number + 1):
            code = _cell_text(cells[code_col]) if code_col < len(cells) else ""
            name = _cell_text(cells[name_col]) if name_col < len(cells) else ""
            if not code and not name:
                continue  # fila en blanco: se ignora sin reportarla
            result.total_rows += 1
            if result.total_rows > MAX_ROWS:
                raise RosterFormatError(f"El Excel supera el máximo de {MAX_ROWS} estudiantes.")

            problem = None
            if not code:
                problem = "Falta el código del estudiante."
            elif not name:
                problem = "Falta el nombre del estudiante."
            elif len(code) > MAX_FIELD_LENGTH or len(name) > MAX_FIELD_LENGTH:
                problem = f"El código y el nombre deben tener como máximo {MAX_FIELD_LENGTH} caracteres."
            elif code in seen:
                problem = f"El código {code!r} ya aparece en la fila {seen[code]}."
            if problem:
                result.errors.append(RowError(current_row, problem))
                continue

            seen[code] = current_row
            result.rows.append(RosterRow(current_row, code, name))

        if result.total_rows == 0:
            raise RosterFormatError("El Excel solo tiene encabezado: no hay estudiantes para cargar.")
        return result
    finally:
        workbook.close()
